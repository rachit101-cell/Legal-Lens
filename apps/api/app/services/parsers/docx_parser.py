"""
LegalLens API — DOCX Parser.

Extracts text and structure from Word documents.
"""

from __future__ import annotations

import hashlib
import io
from typing import BinaryIO

import docx
import structlog

from packages.schemas.domain import CanonicalPage, TextBlock, TextSource

logger = structlog.get_logger()


def extract_docx(
    file_bytes: bytes | BinaryIO,
    document_id: str,
) -> list[CanonicalPage]:
    """
    Parse a DOCX file into canonical pages and text blocks.

    Since DOCX does not have fixed pages in the same way PDF does,
    this creates a single virtual "page" or splits arbitrarily based
    on page breaks if they exist. For MVP, we treat the entire document
    as Page 1, creating a block per paragraph.
    """
    if isinstance(file_bytes, bytes):
        stream = io.BytesIO(file_bytes)
    else:
        stream = file_bytes
        stream.seek(0)

    try:
        doc = docx.Document(stream)
    except Exception as e:
        logger.error("docx_parsing_failed", document_id=document_id, error=str(e))
        raise ValueError(f"Failed to parse DOCX: {e}") from e

    canonical_blocks = []
    page_text_parts = []
    block_order = 0

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        # Create canonical block
        # Bbox is None for DOCX since it's flow layout
        canonical_block = TextBlock(
            block_id=f"{document_id}_p1_b{block_order}",
            text=text,
            bbox=None,
            order=block_order,
        )
        canonical_blocks.append(canonical_block)
        page_text_parts.append(text)
        block_order += 1

    # Handle tables
    for table in doc.tables:
        for row in table.rows:
            row_data = []
            for cell in row.cells:
                text = cell.text.strip()
                if text:
                    row_data.append(text)

            if not row_data:
                continue

            row_text = " | ".join(row_data)

            canonical_block = TextBlock(
                block_id=f"{document_id}_p1_b{block_order}",
                text=row_text,
                bbox=None,
                order=block_order,
            )
            canonical_blocks.append(canonical_block)
            page_text_parts.append(row_text)
            block_order += 1

    full_page_text = "\n\n".join(page_text_parts)
    content_hash = hashlib.sha256(full_page_text.encode("utf-8")).hexdigest()

    # Create a single "page" for the entire DOCX
    canonical_page = CanonicalPage(
        page_id=f"{document_id}_p1",
        page_number=1,
        text=full_page_text,
        text_source=TextSource.NATIVE_TEXT,
        width=None,
        height=None,
        blocks=canonical_blocks,
        content_hash=content_hash,
    )

    return [canonical_page]
